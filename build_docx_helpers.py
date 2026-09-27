import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    """Adds a heading with consistent font and styling."""
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = 'Calibri'
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return h

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    """Adds a body paragraph with standard typography."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def add_code_block(doc, code_text, caption=None):
    """Adds a nicely formatted monospace code block inside a shaded box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>
            <w:left w:val="single" w:sz="16" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    
    if caption:
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_before = Pt(4)
        cap_p.paragraph_format.space_after = Pt(10)
        c_run = cap_p.add_run(caption)
        c_run.font.name = 'Calibri'
        c_run.font.size = Pt(9.5)
        c_run.italic = True
        c_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

def add_centered_image(doc, img_path, width=Inches(5.5), caption=None):
    """Adds a centered image with optional italic caption."""
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=width)
        
        if caption:
            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(10)
            c_run = cap_p.add_run(caption)
            c_run.font.name = 'Calibri'
            c_run.font.size = Pt(9.5)
            c_run.italic = True
            c_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

print("Helper functions defined.")
